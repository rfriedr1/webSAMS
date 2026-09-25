"""Read an uploaded `.xlsx` into a normalized cell grid.

Deliberately thin: the only job here is to turn bytes into a rectangular
list-of-lists of Python values, plus a little sheet-picking smarts. All
interpretation lives in `parser.py`.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

from sams_web.sample_import.vocabulary import (
    SAMPLE_TABLE_SECTION_MARKERS,
    SUBMITTER_SECTION_MARKERS,
    contains_marker,
    normalize_label,
)

# Guard rails for uploads. The real template is ~45 KB and ~1000 rows;
# these limits leave generous headroom while refusing anything that
# would blow up memory.
MAX_UPLOAD_BYTES = 12 * 1024 * 1024
MAX_ROWS = 5_000
MAX_COLS = 80


class WorkbookError(ValueError):
    """Raised when the upload cannot be read as a submission workbook."""


@dataclass(frozen=True)
class Cell:
    """One spreadsheet cell, carrying its address for provenance.

    The review UI shows `address` next to every parsed value so an
    operator can trace a field straight back to the sheet.
    """

    row: int
    col: int
    value: Any

    @property
    def address(self) -> str:
        return f"{get_column_letter(self.col)}{self.row}"

    @property
    def text(self) -> str:
        """The cell rendered as trimmed display text ('' when empty)."""
        if self.value is None:
            return ""
        if isinstance(self.value, datetime):
            return self.value.date().isoformat()
        if isinstance(self.value, date):
            return self.value.isoformat()
        if isinstance(self.value, float) and self.value.is_integer():
            return str(int(self.value))
        return str(self.value).strip()

    @property
    def key(self) -> str:
        """Normalised comparison key for label matching."""
        return normalize_label(self.value)

    @property
    def is_empty(self) -> bool:
        return self.text == ""


@dataclass
class Grid:
    """A single worksheet as a dense grid of `Cell`s (1-indexed access)."""

    sheet_name: str
    rows: list[list[Cell]] = field(default_factory=list)

    @property
    def row_count(self) -> int:
        return len(self.rows)

    @property
    def col_count(self) -> int:
        return max((len(r) for r in self.rows), default=0)

    def cell(self, row: int, col: int) -> Cell:
        """1-indexed cell access; out-of-range yields an empty cell so
        callers never have to bounds-check."""
        if 1 <= row <= len(self.rows):
            line = self.rows[row - 1]
            if 1 <= col <= len(line):
                return line[col - 1]
        return Cell(row=row, col=col, value=None)

    def row(self, row: int) -> list[Cell]:
        if 1 <= row <= len(self.rows):
            return self.rows[row - 1]
        return []

    def row_is_empty(self, row: int) -> bool:
        return all(cell.is_empty for cell in self.row(row))


def _score_sheet(grid: Grid) -> int:
    """Heuristic: how much does this sheet look like a submission sheet?

    The template ships two extra sheets of prose instructions
    ('Ausfüllhilfe', 'Probenhandhabung'); this keeps us off them without
    hard-coding their names, since a customer may rename or delete them.
    """
    score = 0
    for line in grid.rows[:120]:
        for cell in line:
            if cell.is_empty:
                continue
            if contains_marker(cell.value, SUBMITTER_SECTION_MARKERS):
                score += 5
            if contains_marker(cell.value, SAMPLE_TABLE_SECTION_MARKERS):
                score += 5
    # Prose sheets are one narrow column; real sheets are wide.
    if grid.col_count >= 4:
        score += 3
    return score


def load_grids(data: bytes, *, filename: str = "upload.xlsx") -> list[Grid]:
    """Parse workbook bytes into one `Grid` per worksheet."""
    if not data:
        raise WorkbookError("The uploaded file is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise WorkbookError(
            f"File is too large ({len(data) / 1_048_576:.1f} MB). "
            f"The limit is {MAX_UPLOAD_BYTES // 1_048_576} MB."
        )
    if not filename.lower().endswith((".xlsx", ".xlsm")):
        raise WorkbookError(
            "Please upload an Excel workbook (.xlsx). "
            "Legacy .xls files must be re-saved as .xlsx first."
        )

    try:
        # data_only: we want cached formula *results*, not formulas.
        # read_only keeps memory flat on wide sheets.
        workbook = openpyxl.load_workbook(
            io.BytesIO(data), data_only=True, read_only=True
        )
    except Exception as exc:  # noqa: BLE001 - openpyxl raises many types
        raise WorkbookError(
            "That file could not be read as an Excel workbook. "
            "If it was exported from another tool, try re-saving it from Excel."
        ) from exc

    grids: list[Grid] = []
    try:
        for sheet in workbook.worksheets:
            grid = Grid(sheet_name=sheet.title)
            for row_index, row_values in enumerate(
                sheet.iter_rows(max_row=MAX_ROWS, max_col=MAX_COLS, values_only=True),
                start=1,
            ):
                grid.rows.append(
                    [
                        Cell(row=row_index, col=col_index, value=value)
                        for col_index, value in enumerate(row_values, start=1)
                    ]
                )
            # Trim trailing empty rows so row scans terminate quickly.
            while grid.rows and all(cell.is_empty for cell in grid.rows[-1]):
                grid.rows.pop()
            grids.append(grid)
    finally:
        workbook.close()

    if not grids:
        raise WorkbookError("The workbook contains no worksheets.")
    return grids


def pick_submission_grid(grids: list[Grid]) -> Grid:
    """Choose the worksheet that looks most like a submission sheet."""
    non_empty = [g for g in grids if g.row_count > 0]
    if not non_empty:
        raise WorkbookError("Every worksheet in the workbook is empty.")
    return max(non_empty, key=_score_sheet)


__all__ = [
    "MAX_UPLOAD_BYTES",
    "Cell",
    "Grid",
    "WorkbookError",
    "load_grids",
    "pick_submission_grid",
]
